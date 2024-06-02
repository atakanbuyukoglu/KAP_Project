from pathlib import Path
from openpyxl import Workbook, load_workbook
from itertools import islice
from .Metrics import Company
from ..DataFetch.Helpers.RequestWrapper import Request
from ..DataFetch.Helpers.utils import standardize_ticker
from yfinance import Ticker
import numpy as np

class Records():
    
    def __init__(self, file_path, initial_tickers=None, add_tickers=False, online=True) -> None:
        self.file_path = Path(file_path)
        if self.file_path.is_file():
            self.file = load_workbook(filename=self.file_path)
        else:
            self.file = Workbook()
            self.file.save(self.file_path)
        self.tickers = initial_tickers

        # Initialize the file if initial tickers are given
        if add_tickers:
            self.__set_values('Hisse', self.tickers, save=True)

        self.online = online
        self.yahoo_session = Request(sleep_time=1.0)

    def update(self, all=True):
        online = self.online
        self.update_prices(all=all)
        self.update_intrinsic_values(all=all)
        self.online = False
        self.update_intrinsic_values(all=all, quarter=True)
        self.online = online
        #self.file.save(filename=self.file_path)

    def update_ticker(self, ticker: str, save=True):
        ticker = standardize_ticker(ticker)

        sheet = self.file.active
        header_location = self.__get_header_location('Hisse')
        ticker_column = sheet[header_location]
        ticker_row = -1
        for ticker_idx in range(len(ticker_column)):
            if ticker_column[ticker_idx].value == ticker:
                ticker_row = ticker_idx + 1
        if ticker_row == -1:
            raise KeyError(ticker)

        price_header = self.__get_header_location('Fiyat')
        cell_location = price_header + str(ticker_row)
        sheet[cell_location] = self.__get_price(ticker)

        intrinsic_header = self.__get_header_location('İçsel Değer')
        cell_location = intrinsic_header + str(ticker_row)
        sheet[cell_location] = self.__get_intrinsic_value(ticker)
        intrinsic_header = self.__get_header_location('İçsel Değer (Çeyrek)')
        cell_location = intrinsic_header + str(ticker_row)
        sheet[cell_location] = self.__get_intrinsic_value(ticker, quarter=True, online=False)

        if save:
            self.file.save(filename=self.file_path)

    # TODO: Redesägn this function for proper use
    def update_share_counts(self, all=True, save=True, intrinsic_online: bool=None):
        if intrinsic_online is None:
            intrinsic_online = self.online
        # Load the sheet
        sheet = self.file.active
        # Get header location, update the header if needed
        intrinsic_header_str = 'İçsel Değer'
        intrinsic_header_quarter_str = 'İçsel Değer (Çeyrek)'
        intrinsic_header = self.__get_header_location(intrinsic_header_str)
        intrinsic_header_quarter = self.__get_header_location(intrinsic_header_quarter_str)
        ticker_header = self.__get_header_location('Hisse')
        # Update all values on the column
        tickers = self.__get_values(ticker_header)
        values = self.__get_values(intrinsic_header)
        qtr_values = self.__get_values(intrinsic_header_quarter)
        for idx, value in enumerate(values):
            if tickers[idx] is None or tickers[idx] == 'Total':
                continue
            # This happens on empty parts
            if all or value is None:
                try:
                    cell_location = intrinsic_header + str(idx + 2)
                    cell_location_qtr = intrinsic_header_quarter + str(idx + 2)
                    self.__get_share_count(tickers[idx])
                    intr_value = self.__get_intrinsic_value(tickers[idx], quarter=False, online=intrinsic_online)
                    intr_value_qtr = self.__get_intrinsic_value(tickers[idx], quarter=True, online=intrinsic_online)
                    sheet[cell_location] = intr_value
                    sheet[cell_location_qtr] = intr_value_qtr
                except Exception as e:
                    if save:
                        self.file.save(filename=self.file_path)
                    raise e

        if save:
            self.file.save(filename=self.file_path)

    def update_intrinsic_values(self, all=True, quarter=False, save=True):
        intrinsic_header_str = 'İçsel Değer (Çeyrek)' if quarter else 'İçsel Değer'
        self.update_column(intrinsic_header_str, self.__get_intrinsic_value, all=all, save=save, quarter=quarter)

    def update_prices(self, all=True, save=True):
        self.update_column('Fiyat', self.__get_price, all=all, save=save)

    def update_revenue_change(self, all=True, save=True):
        self.update_column('Hasılat Artışı', self.__get_revenue_change, all=all, save=save)

    def update_last_quarter(self, all=True, save=True):
        self.update_column('Son Çeyrek', self.__get_last_quarter, all=all, save=save)

    def update_column(self, target_header_str, target_function, all=True, save=True,  **kwargs):
        # Load the sheet
        sheet = self.file.active
        # Get header location, update the header if needed
        target_header = self.__get_header_location(target_header_str)
        ticker_header = self.__get_header_location('Hisse')
        # Update all values on the column
        tickers = self.__get_values(ticker_header)
        values = self.__get_values(target_header)
        for idx, value in enumerate(values):
            if tickers[idx] is None or tickers[idx] == 'Total':
                continue
            # This happens on empty parts
            if all or value is None:
                cell_location = target_header + str(idx + 2)
                target_value = target_function(tickers[idx],  **kwargs)
                sheet[cell_location] = target_value
        
        if save:
            self.file.save(filename=self.file_path)

    def __get_header_location(self, header_name):
        # Load the sheet
        sheet = self.file.active
        try:
            headers = [cell.value for cell in next(sheet.rows)]
        except StopIteration:
            headers = []
        # Try to get the header
        try:
            header_location = headers.index(header_name)
            header_location = chr(ord('A') + header_location)
        # If it does not work, create the header
        except ValueError:
            # Try to find an empty cell
            header_location = 'A'
            while sheet[header_location + '1'].value:
                header_location = chr(ord(header_location) + 1)
            sheet[header_location + '1'] = header_name
        # Return the header location
        return header_location
    
    def __get_values(self, header_location: str):
        # Load the sheet
        sheet = self.file.active

        header_idx = ord(header_location) - ord('A')
        header_cells = next(islice(sheet.columns, header_idx, None))
        values = [cell.value for cell in header_cells]
        values = values[1:] # Remove the header

        return values

    def __set_values(self, header_name: str, values: list, save=False):
        # Load the sheet
        sheet = self.file.active
        # Get header location
        header_location = self.__get_header_location(header_name)
        # Save the values
        value_idx = 2
        for value in values:
            cell_name = header_location + str(value_idx)
            sheet[cell_name] = value
            value_idx += 1
        if save:
            self.file.save(self.file_path)


    def __get_share_count(self, ticker: str):
        company = Company(ticker, self.file_path.parents[1], update=False)
        print('Obtaining share count of', ticker)
        return company.get_share_count(online=self.online)

    def __get_intrinsic_value(self, ticker: str, quarter=False, online=None):
        if online is None:
            online = self.online
        company = Company(ticker, self.file_path.parents[1], update=online)
        multiplier = 1.0
        return company.default_valuation(quarter=quarter, extra_multiple=multiplier)

    def __get_price(self, ticker: str):
        try:
            stock = Ticker(ticker=ticker + '.IS', session=self.yahoo_session)
            metadata = stock.get_history_metadata()
            price = metadata['regularMarketPrice']
            price = np.round(price, 2)
            print('Price for', ticker, 'obtained:', price)
            return price
        except:
            print('Price for', ticker, 'could not be obtained.')
            return 0.0

    def __get_revenue_change(self, ticker:str, online=None):
        if online is None:
            online = self.online
        company = Company(ticker, self.file_path.parents[1], update=online)
        print('Getting revenue growth for', ticker)
        return company.get_revenue_growth()

    def __get_last_quarter(self, ticker:str, online=None):
        if online is None:
            online = self.online
        company = Company(ticker, self.file_path.parents[1], update=online)
        print('Getting last quarter for', ticker)
        return 'Q' + str(company.last_quarter)
